# -*- coding: utf-8 -*-
import sys, io, os, sqlite3
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

# Accept db path or keywords from command line
# python export_excel.py                          → default db, auto-detect keywords
# python export_excel.py path/to/db.db            → custom db path
# python export_excel.py "keyword1,keyword2"      → default db, custom keywords
arg1 = sys.argv[1] if len(sys.argv) > 1 else ""
if arg1.endswith(".db") and os.path.exists(arg1):
    db = arg1
elif os.path.exists(arg1):
    db = arg1
else:
    db = os.path.join(os.path.dirname(__file__), "database", "sqlite_tables.db")

# If arg1 is not a db path, treat as keywords
cli_keywords = "" if (arg1.endswith(".db") or os.path.exists(arg1)) else arg1

db_name = os.path.splitext(os.path.basename(db))[0]
conn = sqlite3.connect(db)
cur = conn.cursor()

# Build keyword filter from CLI args, or use source_keyword column as fallback
if cli_keywords:
    kw_list = [k.strip() for k in cli_keywords.split(",") if k.strip()]
    kw_conds = " OR ".join([f"title LIKE '%{k}%'" for k in kw_list])
else:
    # Auto-detect: use source_keyword as filter
    cur.execute("SELECT DISTINCT source_keyword FROM douyin_aweme WHERE source_keyword!='' LIMIT 10")
    src_kws = [r[0] for r in cur.fetchall() if r[0]]
    if src_kws:
        kw_conds = " OR ".join([f"source_keyword LIKE '%{k}%' OR title LIKE '%{k}%'" for k in src_kws[:5]])
    else:
        kw_conds = "1=1"  # No filter

# Count and fetch
cur.execute(f"SELECT COUNT(*) FROM douyin_aweme WHERE ({kw_conds}) AND aweme_type != '68'")
total = cur.fetchone()[0]
cur.execute(f"SELECT COUNT(*) FROM douyin_aweme")
raw_total = cur.fetchone()[0]

# If keyword filter returns 0 (e.g. creator mode), export all videos
if total == 0:
    kw_conds = "1=1"
    cur.execute(f"SELECT COUNT(*) FROM douyin_aweme WHERE ({kw_conds}) AND aweme_type != '68'")
    total = cur.fetchone()[0]
    print(f"Raw: {raw_total}, No keyword match, exporting all: {total}")
else:
    print(f"Raw: {raw_total}, Filtered: {total}")

cur.execute(f"""SELECT nickname, follower_count, total_favorited, title, aweme_url, liked_count, comment_count, collected_count, share_count, play_count, create_time, source_keyword, duration
FROM douyin_aweme WHERE ({kw_conds}) AND aweme_type != '68' ORDER BY create_time DESC""")
rows = cur.fetchall()
conn.close()

wb = Workbook()
ws = wb.active
ws.title = "视频数据"

headers = ["序号", "昵称", "粉丝数", "作者获赞", "视频标题", "视频链接", "点赞", "评论", "收藏", "分享", "播放量", "发布时间", "关键词", "视频时长"]
hfont = Font(bold=True, color="FFFFFF")
hfill = Font  # placeholder
from openpyxl.styles import PatternFill, Alignment
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

for i, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = hfont; c.fill = hfill
    c.alignment = Alignment(horizontal='center', wrap_text=True)

for ri, row in enumerate(rows, 2):
    nickname, followers, total_fav, title, url, likes, comments, collects, shares, plays, ctime, kw, dur = row
    ts = ""
    if ctime:
        try: ts = datetime.fromtimestamp(int(ctime)).strftime("%Y-%m-%d %H:%M:%S")
        except: pass

    def ti(v):
        try: return int(v)
        except: return 0

    dur_ms = ti(dur)
    dur_s = f"{dur_ms//1000//60}:{dur_ms//1000%60:02d}" if dur_ms > 0 else ""

    for ci, v in enumerate([ri-1, nickname or "", ti(followers), ti(total_fav), title or "", url or "", ti(likes), ti(comments), ti(collects), ti(shares), ti(plays), ts, kw or "", dur_s], 1):
        cell = ws.cell(row=ri, column=ci, value=v)
        if ci == 5 and url:
            cell.hyperlink = url
            cell.font = Font(color="0563C1", underline="single")

for ci, w in enumerate([6, 16, 10, 10, 50, 45, 10, 10, 10, 10, 10, 20, 14, 10], 1):
    ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w
ws.freeze_panes = 'A2'

out = os.path.join(os.path.expanduser("~"), "Desktop", f"{db_name}_抖音数据.xlsx")
wb.save(out)
print(f"Done: {out} ({len(rows)} rows)")
